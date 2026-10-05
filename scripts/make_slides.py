"""발표자료 10장 (python-pptx). IBDP 양식이 있으면 그 위에 만든다. 수치는 outputs/eval 채점 결과에서.

    python scripts/make_slides.py --template <팀 PPT 양식.pptx>      # 양식 없이 돌리면 빈 16:9 슬라이드에 만듦
"""
import argparse
import os
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "scripts"))
from make_docs import DATA, G_DEMO, G_LANGS, G_TEST, MEMBERS, REPO, SUBTITLE, TEAM, TITLE, YOLO, agg, load_results, pct  # noqa: E402

VIDEO_LEN = os.environ.get("VIDEO_LEN", "2분 40초")
DEMO_SEED = os.environ.get("DEMO_SEED", "5")

FIG = os.path.join(ROOT, "docs")
BLUE = RGBColor(0x28, 0x4E, 0x92)
INK = RGBColor(0x1E, 0x24, 0x30)
GRAY = RGBColor(0x5A, 0x63, 0x70)
RED = RGBColor(0xD9, 0x3A, 0x3A)
GREEN = RGBColor(0x2E, 0x9D, 0x5B)
ORANGE = RGBColor(0xE0, 0x8A, 0x00)
LIGHT = RGBColor(0xEE, 0xF2, 0xF8)
FONT = "맑은 고딕"


def drop_slides(prs):
    lst = prs.slides._sldIdLst
    for sld in list(lst):
        prs.part.drop_rel(sld.rId)
        lst.remove(sld)


def tb(slide, x, y, w, h, lines, size=16, color=INK, bold=False, align=None, anchor=None, fill=None, line=None):
    """lines: 문자열 또는 (문자열, 크기, 색, 굵게) 목록. '**' 로 굵게."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)) if fill is None else \
        slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        if line is None:
            box.line.fill.background()
        else:
            box.line.color.rgb = line
        box.adjustments[0] = 0.08
    tf = box.text_frame
    tf.word_wrap = True
    for side in ("margin_left", "margin_right"):
        setattr(tf, side, Inches(0.12))
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    if anchor or fill is not None:
        tf.vertical_anchor = anchor or MSO_ANCHOR.TOP
    if isinstance(lines, (str, tuple)):
        lines = [lines]
    for i, ln in enumerate(lines):
        if isinstance(ln, tuple):
            text, sz, col, bd = (list(ln) + [None, None, None])[:4]
        else:
            text, sz, col, bd = ln, None, None, None
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        p.space_after = Pt(4)
        for k, part in enumerate(text.split("**")):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            r.font.name = FONT
            r.font.size = Pt(sz or size)
            r.font.bold = bool(bd if bd is not None else bold) or k % 2 == 1
            r.font.color.rgb = col or color
    return box


def pic(slide, path, x, y, w=None, h=None):
    if not os.path.exists(path):
        return tb(slide, x, y, w or 3, 0.4, f"[그림 없음: {os.path.basename(path)}]", 10, GRAY)
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)


def table(slide, rows, x, y, w, col_w, size=12, header_fill=BLUE, row_h=0.36, bold_last=False):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    t = shp.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = t.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.06)
            c.margin_top = c.margin_bottom = Inches(0.03)
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER if j > 0 or i == 0 else PP_ALIGN.LEFT
            for k, part in enumerate(str(val).split("**")):
                if not part:
                    continue
                r = p.add_run()
                r.text = part
                r.font.name = FONT
                r.font.size = Pt(size)
                r.font.bold = i == 0 or k % 2 == 1 or (bold_last and i == len(rows) - 1)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if i == 0 else INK
            c.fill.solid()
            c.fill.fore_color.rgb = header_fill if i == 0 else (LIGHT if i % 2 == 0 else RGBColor(0xFF, 0xFF, 0xFF))
    return shp


def content_slide(prs, title, layout_idx):
    """제목 길이가 들쭉날쭉해 기본 글자 크기로는 2줄로 접히며 바로 아래 내용과 겹치는 제목이 있어(특히 영문 용어를
    괄호로 덧붙인 제목), 제목 글자 크기를 고정해 최대 2줄 높이를 예측 가능하게 만든다. 각 슬라이드의 본문은 이
    2줄 높이(약 1.5in)보다 아래에서 시작해야 겹치지 않는다."""
    s = prs.slides.add_slide(prs.slide_layouts[layout_idx])
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text = title
            for p in ph.text_frame.paragraphs:
                for r in p.runs:
                    r.font.name = FONT
                    r.font.size = Pt(28)
        elif idx not in (12, 10, 11):
            ph._element.getparent().remove(ph._element)
    return s


def build(template, out):
    rs = load_results()
    A = agg(rs)
    a, b, rc = A["a"], A["b"], A["rc"]
    nh, ns = a["hazard_total"], a["safe_total"]
    if template and os.path.exists(template):
        prs = Presentation(template)
        drop_slides(prs)
        L_TITLE, L_BODY = 0, 1
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        L_TITLE, L_BODY = 0, 5

    # 1. 표지
    s = prs.slides.add_slide(prs.slide_layouts[L_TITLE])
    for ph in list(s.placeholders):
        idx = ph.placeholder_format.idx
        if idx == 0:
            ph.text = TITLE
        elif idx == 1:
            ph.text = f"{SUBTITLE}\n\n{TEAM} · {MEMBERS}\n2026-10-06"
            for k, p in enumerate(ph.text_frame.paragraphs):
                for r in p.runs:
                    r.font.size = Pt(16 if k == 0 else 18)
                    r.font.name = FONT
    tb(s, 2.5, 1.15, 8.3, 0.45, "제4회 경남AI·SW경진대회 · ③ 제조·피지컬 AI Agent", 16, GRAY, align=PP_ALIGN.CENTER)

    # 2. 문제
    s = content_slide(prs, "1. 문제: 눈으로 보고, 한국어로만 안내한다", L_BODY)
    # 제목이 길어 2줄로 접히는 경우가 있어(한글 제목 폭이 커 자동 줄바꿈), 아래 내용은 1.8in 부터 시작해 2줄 제목과도 안 겹치게 함
    tb(s, 0.6, 1.8, 4.2, 2.3, [("25.2%", 66, RED, True), ("2025년 사고재해자 113,305명 중 넘어짐 28,606명, 재해 유형 1위", 14, GRAY),
                                ("(고용노동부 「2025년 산업재해 현황」)", 11, GRAY)], align=PP_ALIGN.CENTER)
    tb(s, 0.6, 4.5, 4.2, 2.6, [("공장·창고의 위험 요소", 18, BLUE, True), "• 바닥 유출 → 미끄러짐", "• 통로에 방치된 공구·자재 → 걸려 넘어짐",
                                "• 불안정 적재 → 무너짐·맞음", "• 쓰러지거나 가로막힌 소화기 → 화재 초기 대응 실패"], 15)
    tb(s, 5.3, 1.8, 7.5, 0.5, ("지금의 현장 점검", 20, BLUE, True))
    for k, (h, body) in enumerate([("놓치기 쉽다", "작업자가 걸으며 눈으로만 확인. 잠깐 보고 지나친 것, 경로에서 먼 것은 빠짐"),
                                   ("기록이 남지 않는다", "무엇을 어디서 어떤 근거로 위험하다고 봤는지 남지 않아 조치 지시와 추적이 어려움"),
                                   ("다가가는 순간을 모른다", "작업자가 위험물 바로 옆을 지나가는 아차사고를 관리자가 알 수 없음"),
                                   ("한국어로만 안내된다", "외국인 근로자는 표지판·방송을 알아듣지 못해 위험을 더 늦게 알아차림")]):
        tb(s, 5.3, 2.4 + k * 1.25, 7.5, 1.1, [(h, 18, INK, True), (body, 13, GRAY)], fill=LIGHT, anchor=MSO_ANCHOR.MIDDLE)

    # 3. 해결 방안
    X = A["x"]
    s = content_slide(prs, "2. 해결: 판정·경고하고 작업자 언어로 답하는 AI 에이전트", L_BODY)
    # 제목이 2줄로 접혀 아래 내용과 안 겹치게 전체를 조금 내려서 배치 (1. 문제 슬라이드와 같은 이유)
    tb(s, 0.6, 1.65, 12.1, 0.9, [("바디캠 하나로 위험/안전과 **공구 이름**을 판정하고, 애매한 것은 **'주의'로 분류**해 바디캠이 다시 지나칠 때 **재판단**한다. "
                                 "**위험 영역**에 닿기 직전이면 **음성으로 경고**하고, 작업자가 **손가락 1~3개**를 보이면 **작업자 언어**로 안내·호출·SOS 를 처리한다", 15, INK)])
    pic(s, os.path.join(FIG, "workflow.png"), 0.6, 2.6, w=12.1)
    feats = [("① 판정", "위험/안전 상태\n공구 이름 8종", BLUE), ("② 재관측 재판단", "'주의' 분류 →\n재관측 시 LLM 판단", ORANGE),
             ("③ 위험 영역", "라바콘·DANGER 표지\n스스로 판단한 주변", RED), ("④ 음성 경고", "위험 \"멈추세요\"\n주의 \"발밑 확인\"", RGBColor(0x9B, 0x4D, 0xCA)),
             ("⑤ 손동작 명령", "손가락 1~3 → LLM 판단\n중·영·일 음성 안내",RGBColor(0x0A, 0x7C, 0xC4)), ("⑥ 조치 지시서", "우선순위·위치\n영역 지도와 기록", GREEN)]
    for k, (h, body, col) in enumerate(feats):
        tb(s, 0.6 + k * 2.04, 4.85, 1.94, 1.75, [(h, 15, col, True), (body, 12, INK)], fill=LIGHT, align=PP_ALIGN.CENTER)
    tb(s, 0.6, 6.75, 12.1, 0.4, "로봇 없이 작업자 바디캠(스피커)만으로 동작, 별도 CCTV 설비 불필요 · 위험 데이터는 디지털 트윈에서 합성", 13, GRAY, align=PP_ALIGN.CENTER)

    # 4. 구조
    s = content_slide(prs, "3. 에이전트 구조 (Goal · Plan · Reasoning · Tool · Memory · Feedback)", L_BODY)
    # 제목이 2줄이라(28pt) 그림을 조금 내리고 작게(11.3in) 줄여 아래 표와 안 겹치게 함
    pic(s, os.path.join(FIG, "architecture.png"), 1.0, 1.5, w=11.3)
    table(s, [["Goal", "Planning", "Reasoning", "Tool Use", "Memory", "Feedback"],
              ["판정, 위험 영역, 경고, 작업자 요청 응답, 조치 지시서", "도면으로 점검표 8곳 계산", "애매한 판정 2단계 분류, 공구 자리, 영역 설정, 경고 시점·톤, 손동작 → 명령",
               "바디캠, 손 인식, LLM(Qwen2.5-7B)이 부르는 도구(매뉴얼 RAG 포함), 음성", "위험물 대장, 위험 영역, 경고·요청 기록", "주의 물체 재관측 재판단, 현장 확인 권고"]],
          0.5, 5.9, 12.3, [2.05] * 6, size=11, row_h=0.5)

    # 5. 디지털 트윈과 데이터
    s = content_slide(prs, "4. 디지털 트윈과 학습 데이터", L_BODY)
    pic(s, os.path.join(FIG, "fig_zone.jpg"), 0.5, 1.3, w=6.0)
    tb(s, 0.5, 4.75, 6.0, 1.9, [("NVIDIA Isaac Sim 6.0 실사 창고", 16, BLUE, True), "• NVIDIA 소품 (소화기, 표지판, 라바콘, 상자, 팔레트, 작업자)",
                                 "• Poly Haven 실물 스캔 공구 (CC0) + 직접 만든 전동톱, DANGER 표지", "• Replicator 로 정답 박스 자동 생성, 정답표로 채점"], 13)
    table(s, [["종류", "위험", "안전"], ["바닥 유출", "조치 없음", "표지판·라바콘"], ["공구 8종", "통로 바닥", "작업대 위"],
              ["적재", "기울어짐", "반듯함"], ["소화기", "쓰러짐 / 가로막힘", "제자리"], ["영역 표시", "라바콘 링, DANGER 표지", ""]],
          6.9, 1.3, 5.9, [1.7, 2.2, 2.0], size=12, row_h=0.34)
    tb(s, 6.9, 3.55, 5.9, 3.1, [(f"합성 데이터 {DATA['n']}장 → YOLO26s 19 클래스", 16, BLUE, True),
                                 f"• {DATA['desc']}", "• 바디캠 70%·자유 시점 30%, 조명·흔들림 무작위",
                                 "• 같은 종류를 위험/안전 상태로 함께 배치 → 상태를 학습",
                                 "• 공구 8종·운반 카트는 종류로, 공구 위험/안전은 놓인 자리로",
                                 f"• 검증 mAP50 **{YOLO['map50']}** (mAP50-95 {YOLO['map5095']})"], 13)

    # 6. 핵심 구현 1
    s = content_slide(prs, "5. 핵심 구현 ①  판정, 위험 영역, 음성 경고", L_BODY)
    pic(s, os.path.join(FIG, "fig_body.jpg"), 0.5, 1.3, w=6.1)
    pic(s, os.path.join(FIG, "fig_zone_map.png"), 6.75, 1.3, w=6.1)
    tb(s, 0.5, 4.85, 6.1, 1.9, [("바디캠 판정 + 공구 이름", 16, BLUE, True), "• ByteTrack 추적으로 같은 물체 판정을 모아 3프레임이면 확정",
                                 "• 공구는 망치·삽 등 이름까지, 놓인 자리로 위험(바닥)/안전(작업대)", "• 박스를 바닥에 투영해 위치 → 위험물 대장"], 13)
    tb(s, 6.75, 4.85, 6.1, 1.9, [("위험 영역과 음성 경고", 16, RED, True), "• 라바콘 2개 이상 묶음, DANGER 표지, 유출·불안정 적재 주변 → 영역 (소화기류 제외)",
                                  "• 회피형 위험 1.5 m·영역 0.5 m 안 → 위험 \"멈추세요!\", 주의 \"발밑을 확인하세요\"",
                                  "• 2 m 안은 경고 없이 근접만 기록, 오늘 작업 대상 물체는 경고 제외"], 13)

    # 7. 핵심 구현 2
    s = content_slide(prs, "6. 핵심 구현 ②  애매한 판정 2단계 분류와 재관측 재판단", L_BODY)
    # 제목이 2줄이라 전체를 조금 내려 배치 (1. 문제 슬라이드와 같은 이유)
    pic(s, os.path.join(FIG, "fig_rejudge.jpg"), 0.5, 1.6, w=5.6)
    tb(s, 0.5, 4.8, 5.6, 2.3, [("즉시 분류 (규칙)", 15, ORANGE, True), "• 바디캠에 잠깐 보이거나 판정이 엇갈린 물체",
                                "• 후보 중 고위험 클래스가 하나라도 있으면 즉시 위험으로 확정", "• 전부 저위험이면 '주의'로 분류 (안전·위험 미확정)",
                                ("재관측 재판단", 15, ORANGE, True), "• 주의 물체를 바디캠이 같은 자리에서 다시 포착하면 트리거",
                                "• 로컬 LLM(LangGraph)이 이전+새 증거로 재판단, 실패하면 규칙 폴백", "• 끝까지 안 보이면 주의 유지 → 조치 지시서에 현장 확인 권고"], 12)
    pic(s, os.path.join(FIG, "fig_dashboard_crop.png"), 6.4, 1.6, w=6.45)
    tb(s, 6.4, 6.65, 6.45, 0.45, "조치 지시서: 평면도(번호 = 우선순위, 빨간 영역 = 위험 영역), 조치 목록, 음성 경고 기록", 12, GRAY, align=PP_ALIGN.CENTER)

    # 8. 손동작 명령 (시연 영상은 따로 제출)
    s = content_slide(prs, "7. 핵심 구현 ③  손동작 명령 → 작업자 언어 안내", L_BODY)
    # fig_gestures.jpg 가 이제 3칸(손가락 1~3)이라 5칸일 때보다 가로세로 비율이 달라져, 너비 대신 높이로 맞춰 아래 내용과 안 겹치게 함
    pic(s, os.path.join(FIG, "fig_gestures.jpg"), 0.5, 1.2, h=1.9)
    pic(s, os.path.join(FIG, "fig_gesture.jpg"), 0.5, 3.75, w=5.0)
    tb(s, 5.75, 3.7, 3.6, 2.95, [("어떻게 처리하나", 15, BLUE, True), "• MediaPipe 손 관절 21점 → 손가락 1~3, 멈춘 손에서 3번 연속이면 명령",
                                  "• LangGraph + 로컬 Qwen2.5-7B 가 도구 (화면 물체, 장비 정보, 위험물 대장, TBM) 로 보고 무엇을 말할지·근거 결정",
                                  "• 말은 검수한 4개 언어 문장 틀 (번역기는 '안전화→seat belt' 처럼 틀림)"], 11, fill=LIGHT)
    tb(s, 9.55, 3.7, 3.25, 2.95, [("결과", 15, BLUE, True), (f"{G_TEST['ok']}/{G_TEST['trials']}", 30, INK, True),
                                   (f"경로 8곳·조명 바꿔 손가락 1~3, 다른 명령 {G_TEST['wrong']}번, 움직이는 손 잘못 실행 0", 11, GRAY),
                                   (f"시연 {G_DEMO['recognized']}/{G_DEMO['shown']} · {', '.join(G_LANGS) or '중·영·일'} · LLM 결정 {G_DEMO.get('llm', 0)}/{G_DEMO.get('requests', 0)}", 12, INK)], fill=LIGHT)
    tb(s, 0.5, 6.7, 12.3, 0.35, f"시연 영상 ({VIDEO_LEN}): TBM(2) → 순찰·판정 → 운반 카트 설명(1) → 상자 싣고 끌기 → 관리자 호출·SOS(3) → 조치 지시서",
       12, GRAY, align=PP_ALIGN.CENTER)

    # 9. 성과
    s = content_slide(prs, f"8. 성과: 학습에 안 쓴 시나리오 {A['n']}개, 정답표 채점", L_BODY)
    table(s, [["창고 전체 물체 기준", "바디캠 프레임 원시판정", "에이전트"],
              [f"위험 물체를 위험으로 ({nh}개)", f"{b['hazard_found']} ({pct(b['hazard_found'], nh)})", f"**{a['hazard_found']} ({pct(a['hazard_found'], nh)})**"],
              [f"안전 물체를 안전으로 ({ns}개)", f"{b['safe_ok']} ({pct(b['safe_ok'], ns)})", f"**{a['safe_ok']} ({pct(a['safe_ok'], ns)})**"],
              ["위험↔안전 거꾸로 판정", b["hazard_as_safe"] + b["safe_as_hazard"], a["hazard_as_safe"] + a["safe_as_hazard"]],
              ["없는 위험 판정 (박스 수 / 보고 건수)", f"{b['false_hazard_boxes']}개", f"{a['false_reports']}건"],
              ["현장 확인 요청 (실제 물체)", "-", f"{a['need_check']} ({a['need_check_real']})"]],
          0.5, 1.35, 7.4, [3.6, 1.9, 1.9], size=14, row_h=0.48)
    table(s, [["새 기능", "결과"],
              ["위험 영역 (라바콘·DANGER 표지) 인식", f"**{X['zone_found']}/{X['zone_gt']}** ({pct(X['zone_found'], X['zone_gt'])})"],
              ["에이전트가 판단한 위험 영역 (실제 위험 주변)", f"{X['zone_agent']} ({X['zone_agent_real']})"],
              ["닿기 직전 사건에 음성 경고", f"**{X['voice_warned']}/{X['voice_events']}** ({pct(X['voice_warned'], X['voice_events'])})"],
              ["공구 이름 맞힘", f"**{X['tool_named']}/{X['tool_total']}** ({pct(X['tool_named'], X['tool_total'])})"],
              ["손동작 명령 인식 (손가락 1~3)", f"**{G_TEST['ok']}/{G_TEST['trials']}** ({pct(G_TEST['ok'], G_TEST['trials'])})"]],
          0.5, 4.3, 7.4, [4.6, 2.8], size=13, row_h=0.42)
    tb(s, 8.3, 1.35, 4.5, 2.6, [("애매한 판정과 재판단", 15, BLUE, True),
                                 f"애매한 판정 {rc['ambiguous_hazard'] + rc['ambiguous_caution']}건: 즉시 위험 확정 {rc['ambiguous_hazard']}, 주의 분류 {rc['ambiguous_caution']}",
                                 f"주의 재관측 재판단 {rc['caution_rejudged']}건 중 확정 전환 {rc['caution_rejudged_confirmed']} (맞게 고침 {rc['changed_correct']})",
                                 "끝까지 재관측 안 되면 '주의' 유지 → 현장 확인 권고"], 13, fill=LIGHT)
    tb(s, 8.3, 4.2, 4.5, 2.5, [("점검표", 15, RED, True), (f"{a['need_check']} 건", 30, INK, True),
                                (f"현장 확인 요청 (실제 물체 {a['need_check_real']}건)", 12, GRAY)], fill=LIGHT)

    # 10. 한계·발전·BM
    from collections import Counter
    miss = Counter(row["class"] for r in rs for row in r["agent"]["evaluation"]["rows"] if row["hazard"] and row["after"] is None)
    s = content_slide(prs, "9. 한계, 발전 계획, 비즈니스 모델", L_BODY)
    tb(s, 0.5, 1.3, 4.0, 5.4, [("한계 (현장 적용 전 보완점)", 16, RED, True), "• 합성 데이터만 사용, 실사 미검증",
                                f"• 작은 공구를 놓침 (못 찾은 위험 {sum(miss.values())}개 중 {miss.get('tool_floor', 0)}개가 바닥 공구)",
                                "• Isaac 6.0.1 torch 가 CPU 전용이라 시뮬 안 YOLO 는 CPU (실배치는 GPU·엣지)",
                                "• 영역 크기·경고 거리·즉시 분류 기준은 규칙 (재판단만 LLM)", "• 손동작은 시뮬 맨손만, 다국어 음성은 온라인 TTS",
                                "• 움직이는 위험(지게차)은 위치추적 태그(RTLS/UWB) 개념으로 접근만 감지, 영상 인식은 아직 없음"], 13, fill=LIGHT)
    tb(s, 4.7, 1.3, 4.0, 5.4, [("발전 계획", 16, BLUE, True), ("1단계 실사 검증", 14, INK, True), "실습실·창고 사진으로 미세조정, 실제 라바콘·표지",
                                ("2단계 파일럿 (Level 4)", 14, INK, True), "바디캠(스피커) 현장 시험, 재관측 재판단 기준 보정",
                                ("3단계 연동", 14, INK, True), "MES·작업지시 연동, 조치 완료 피드백으로 기준 보정, 지게차 영상 인식 추가"], 13, fill=LIGHT)
    tb(s, 8.9, 1.3, 3.95, 5.4, [("비즈니스 모델", 16, GREEN, True), ("대상", 14, INK, True), "중소 물류창고·제조공장 (사고재해 40%가 5~49인 사업장)",
                                 ("제공", 14, INK, True), "작업자 바디캠에 붙이는 소프트웨어 구독, 별도 CCTV 설비 불필요 (사업장·인원 수 기준)",
                                 ("차별점", 14, INK, True), "현장 표시(라바콘·표지)를 그대로 위험 영역으로 읽고, 디지털 트윈으로 그 현장 데이터를 만들어 학습",
                                 ("기대효과", 14, INK, True), "점검 누락·기록 공백 감소, 닿기 전 경고로 아차사고 예방"], 13, fill=LIGHT)
    tb(s, 0.5, 6.75, 12.3, 0.3, f"소스코드 {REPO}", 11, GRAY, align=PP_ALIGN.CENTER)
    prs.save(out)
    return out


def main():
    p = argparse.ArgumentParser()
    team_tpl = os.path.join(os.path.dirname(ROOT), "조현준_IBDPppt양식.pptx")      # 팀 양식 (저장소 밖, 있으면 기본으로)
    p.add_argument("--template", default=os.environ.get("SLIDE_TEMPLATE") or (team_tpl if os.path.exists(team_tpl) else None),
                   help="제목/제목+내용 레이아웃이 있는 pptx 양식")
    p.add_argument("--out", default=os.path.join(ROOT, "submission", "발표자료.pptx"))
    a = p.parse_args()
    os.makedirs(os.path.dirname(a.out), exist_ok=True)
    print(build(a.template, a.out))


if __name__ == "__main__":
    main()
